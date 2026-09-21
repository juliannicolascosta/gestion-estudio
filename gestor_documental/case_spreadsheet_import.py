"""Read, preview and import cases from ordinary XLSX/CSV spreadsheets."""

from __future__ import annotations

import csv
import re
import shutil
import unicodedata
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Callable, Iterable

from .case_data import build_case_caption, ensure_system_metadata
from .models import Case
from .services import create_case, list_cases, read_case_metadata, save_case_metadata


FIELD_LABELS = {
    "actor": "Actor",
    "defendant": "Demandado",
    "cause": "Causa",
    "case_number": "N° Expediente",
    "notes": "Observaciones",
    "original_caption": "Carátula original",
}

HEADER_ALIASES = {
    "actor": {"actor"},
    "defendant": {"demandado"},
    "cause": {"causa"},
    "case_number": {"n expediente", "numero expediente", "expediente"},
    "notes": {"observaciones"},
    "original_caption": {"caratula original", "caratula"},
}


class ImportStatus(str, Enum):
    NEW = "NUEVO"
    POSSIBLE_DUPLICATE = "POSIBLE DUPLICADO"
    EXISTS = "YA EXISTE"
    REVIEW = "REQUIERE REVISIÓN"


@dataclass
class SpreadsheetData:
    headers: list[str]
    rows: list[list[str]]
    mapping: dict[str, int]
    ambiguous: set[str]


@dataclass
class ImportRow:
    source_row: int
    actor: str = ""
    defendant: str = ""
    cause: str = ""
    case_number: str = ""
    notes: str = ""
    original_caption: str = ""
    status: ImportStatus = ImportStatus.NEW
    reason: str = ""
    selected: bool = True


@dataclass
class ImportOutcome:
    row: ImportRow
    case: Case | None = None
    error: str = ""


def _plain(value: object) -> str:
    return " ".join(str(value or "").split()).strip()


def _comparison(value: str) -> str:
    return _plain(value).casefold()


def _header_key(value: str) -> str:
    text = unicodedata.normalize("NFKD", _plain(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold().replace("º", " ").replace("°", " ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def detect_mapping(headers: list[str]) -> tuple[dict[str, int], set[str]]:
    matches: dict[str, list[int]] = {field: [] for field in HEADER_ALIASES}
    for index, header in enumerate(headers):
        key = _header_key(header)
        for field, aliases in HEADER_ALIASES.items():
            if key in aliases:
                matches[field].append(index)
    mapping = {field: indexes[0] for field, indexes in matches.items() if len(indexes) == 1}
    ambiguous = {field for field, indexes in matches.items() if len(indexes) > 1}
    return mapping, ambiguous


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    raw = path.read_text(encoding="utf-8-sig")
    sample = raw[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    values = list(csv.reader(raw.splitlines(), delimiter=delimiter))
    return _split_table(values)


def _read_xlsx(path: Path) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Falta el componente para leer archivos XLSX.") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = [[_plain(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        return _split_table(values)
    finally:
        workbook.close()


def _split_table(values: list[list[object]]) -> tuple[list[str], list[list[str]]]:
    nonempty = [row for row in values if any(_plain(value) for value in row)]
    if not nonempty:
        raise ValueError("El archivo no contiene filas para importar.")
    headers = [_plain(value) for value in nonempty[0]]
    if not any(headers):
        raise ValueError("El archivo no contiene encabezados válidos.")
    rows = [[_plain(value) for value in row] for row in nonempty[1:]]
    return headers, rows


def read_spreadsheet(path: Path) -> SpreadsheetData:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        headers, rows = _read_csv(path)
    elif suffix == ".xlsx":
        headers, rows = _read_xlsx(path)
    else:
        raise ValueError("Elegí un archivo XLSX o CSV.")
    mapping, ambiguous = detect_mapping(headers)
    return SpreadsheetData(headers, rows, mapping, ambiguous)


def mapped_rows(data: SpreadsheetData, mapping: dict[str, int]) -> list[ImportRow]:
    if "actor" not in mapping:
        raise ValueError("Indicá qué columna contiene el actor.")

    def value(row: list[str], field: str) -> str:
        index = mapping.get(field)
        return row[index] if index is not None and index < len(row) else ""

    result = []
    for source_row, row in enumerate(data.rows, start=2):
        if not any(row):
            continue
        item = ImportRow(
            source_row=source_row,
            actor=value(row, "actor"),
            defendant=value(row, "defendant"),
            cause=value(row, "cause"),
            case_number=value(row, "case_number"),
            notes=value(row, "notes"),
            original_caption=value(row, "original_caption"),
        )
        if not item.actor:
            item.status = ImportStatus.REVIEW
            item.reason = "Falta el actor"
            item.selected = False
        result.append(item)
    if not result:
        raise ValueError("El archivo no contiene casos para importar.")
    return result


def classify_rows(rows: list[ImportRow], existing_cases: Iterable[Case]) -> list[ImportRow]:
    numbers: set[str] = set()
    descriptions: set[tuple[str, str, str]] = set()
    for case in existing_cases:
        metadata = read_case_metadata(case)
        number = _comparison(metadata.get("CUIJ", "") or metadata.get("Número de expediente", ""))
        if number:
            numbers.add(number)
        descriptions.add((
            _comparison(metadata.get("Actor", "") or metadata.get("Nombre completo", "")),
            _comparison(metadata.get("Demandado", "") or metadata.get("Empleador principal", "")),
            _comparison(metadata.get("Causa", "") or metadata.get("Concepto del reclamo RAEO", "")),
        ))

    seen_numbers = set(numbers)
    seen_descriptions = set(descriptions)
    for row in rows:
        if not row.actor:
            continue
        number = _comparison(row.case_number)
        description = (
            _comparison(row.actor), _comparison(row.defendant), _comparison(row.cause),
        )
        if number and number in seen_numbers:
            row.status = ImportStatus.EXISTS
            row.reason = "Coincide el número de expediente"
            row.selected = False
        elif not number and description in seen_descriptions:
            row.status = ImportStatus.POSSIBLE_DUPLICATE
            row.reason = "Coinciden actor, demandado y causa"
            row.selected = False
        else:
            row.status = ImportStatus.NEW
            row.reason = "Sin Nº de expediente" if not number else ""
            row.selected = True
        if number:
            seen_numbers.add(number)
        seen_descriptions.add(description)
    return rows


def preview_spreadsheet(path: Path, existing_cases: Iterable[Case], mapping: dict[str, int] | None = None) -> tuple[SpreadsheetData, list[ImportRow]]:
    data = read_spreadsheet(path)
    effective_mapping = mapping if mapping is not None else data.mapping
    return data, classify_rows(mapped_rows(data, effective_mapping), existing_cases)


def row_metadata(row: ImportRow, professional: str = "") -> dict[str, str]:
    metadata = {
        "Actor": row.actor,
        "Demandado": row.defendant,
        "Causa": row.cause,
        "CUIJ": row.case_number,
        "Observaciones": row.notes,
        "Carátula original": row.original_caption,
    }
    return ensure_system_metadata(
        {key: value for key, value in metadata.items() if value},
        professional=professional,
    )


def import_rows(
    study_root: Path,
    rows: Iterable[ImportRow],
    *,
    professional: str = "",
    progress: Callable[[int, int], None] | None = None,
) -> list[ImportOutcome]:
    selected = [row for row in rows if row.selected]
    outcomes: list[ImportOutcome] = []
    for index, row in enumerate(selected, start=1):
        case = None
        try:
            metadata = row_metadata(row, professional)
            folder_name = build_case_caption(metadata, row.actor)
            case = create_case(study_root, folder_name)
            save_case_metadata(case, metadata)
            outcomes.append(ImportOutcome(row, case=case))
        except Exception as error:
            cleanup_error = ""
            if case is not None and case.path.is_dir():
                try:
                    shutil.rmtree(case.path)
                except OSError as cleanup:
                    cleanup_error = f"; no se pudo retirar la carpeta incompleta: {cleanup}"
            outcomes.append(ImportOutcome(row, error=f"{error}{cleanup_error}"))
        if progress:
            progress(index, len(selected))
    return outcomes


def existing_cases(study_root: Path) -> list[Case]:
    return list_cases(study_root)
